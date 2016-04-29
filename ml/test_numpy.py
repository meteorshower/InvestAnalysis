#coding=utf8

import numpy as np
from docx import Document
from docx.table import  Table

def linalg():
    a = np.array([[5,-1,2], [-2,6,9], [-7,5,-3]])
    print a
    b = np.array([7,0,-7])
    print b
    x = np.linalg.solve(a,b)
    print x
    print np.dot(a,x)

if __name__ == "__main__":
    d = Document('20160509.docx')
    for tab in d.tables:
        t = tab.cell(0,0).text
        pass
    for sec in d.sections:
        print sec.start_type
    pass
